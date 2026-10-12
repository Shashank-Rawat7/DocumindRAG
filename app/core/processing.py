import logging
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core import storage
from app.core.database import SessionLocal
from app.models.document import Document

logger = logging.getLogger("documind")


class ExtractionError(Exception):
    """A failure whose message is safe to show to the user."""


def extract_text(path: Path) -> str:
    if not path.is_file():
        raise ExtractionError("Stored file not found")

    ext = path.suffix.lower()

    if ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8-sig", errors="replace")
    elif ext == ".pdf":
        reader = PdfReader(path)
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif ext == ".docx":
        doc = DocxDocument(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(c for c in cells if c))
        text = "\n".join(parts)
    else:
        raise ExtractionError("Unsupported file type")

    text = text.replace("\x00", "")
    if not text.strip():
        raise ExtractionError("No extractable text found (the file may be empty or scanned)")

    return text


def _mark_failed(db, document_id: int, reason: str) -> None:
    doc = db.get(Document, document_id)
    if doc is None:
        return
    doc.processing_status = "failed"
    doc.processing_error = reason[:500]
    db.commit()


def _run(db, document_id: int) -> None:
    doc = db.get(Document, document_id)
    if doc is None or not doc.stored_name:
        return

    doc.processing_status = "processing"
    doc.processing_error = None
    db.commit()

    try:
        text = extract_text(storage.get_file_path(doc.stored_name))
        doc.extracted_text = text
        doc.processing_status = "completed"
        db.commit()
    except ExtractionError as exc:
        db.rollback()
        _mark_failed(db, document_id, str(exc))
    except Exception:
        db.rollback()
        logger.error(
            "Document processing failed | document_id=%s", document_id, exc_info=True
        )
        _mark_failed(db, document_id, "Failed to extract text from this file")


def process_document(document_id: int) -> None:
    db = SessionLocal()
    try:
        _run(db, document_id)
    except Exception:
        logger.error(
            "Document processing crashed | document_id=%s", document_id, exc_info=True
        )
        try:
            db.rollback()
            _mark_failed(db, document_id, "Processing failed unexpectedly")
        except Exception:
            logger.error(
                "Could not record failure | document_id=%s", document_id, exc_info=True
            )
    finally:
        db.close()