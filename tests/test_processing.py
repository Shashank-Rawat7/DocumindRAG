import io
import logging
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document as DocxDocument
from pypdf import PdfReader, PdfWriter

from app.core import processing
from app.core.database import SessionLocal
from app.core.processing import ExtractionError, extract_text
from app.models.document import Document
from app.models.user import User

DOCUMENTS = "/api/v1/documents"
UPLOAD = f"{DOCUMENTS}/upload"
SAMPLE_PDF = Path(__file__).parent / "fixtures" / "sample.pdf"

GENERIC_FAILURE = "Failed to extract text from this file"


# ---------- helpers ----------

@pytest.fixture
def sample_pdf():
    assert SAMPLE_PDF.is_file(), f"Missing test fixture: copy sample.pdf to {SAMPLE_PDF}"
    return SAMPLE_PDF


def docx_bytes(paragraphs=(), table=None):
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def blank_pdf_bytes():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def two_page_pdf_bytes(sample_pdf):
    writer = PdfWriter()
    for _ in range(2):
        writer.append(PdfReader(sample_pdf))
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def write(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    return path


def put_file(tmp_path, name, data):
    write(tmp_path, name, data)
    return name


def make_document(db, owner_id, *, stored_name=None, status="pending", error=None):
    doc = Document(
        filename="file.txt",
        owner_id=owner_id,
        stored_name=stored_name,
        processing_status=status,
        processing_error=error,
    )
    db.add(doc)
    db.commit()
    return doc.id


def fresh(db, doc_id):
    db.expire_all()
    return db.get(Document, doc_id)


def upload(client, headers, name, content):
    return client.post(
        UPLOAD,
        files={"file": (name, content, "application/octet-stream")},
        headers=headers,
    )


def explode(*args, **kwargs):
    raise RuntimeError("boom")


@pytest.fixture
def owner_id(make_user, db):
    make_user()
    return db.query(User).one().id


# ---------- extract_text: plain text ----------

@pytest.mark.parametrize("name", ["a.txt", "a.md"])
def test_extract_text_reads_plain_text_files(tmp_path, name):
    path = write(tmp_path, name, b"Hello\nWorld")

    assert extract_text(path) == "Hello\nWorld"


def test_extract_text_handles_uppercase_extensions(tmp_path):
    assert extract_text(write(tmp_path, "UPPER.TXT", b"upper")) == "upper"


def test_extract_text_strips_a_utf8_byte_order_mark(tmp_path):
    assert extract_text(write(tmp_path, "bom.txt", b"\xef\xbb\xbfhello")) == "hello"


def test_extract_text_replaces_invalid_utf8_instead_of_failing(tmp_path):
    assert extract_text(write(tmp_path, "latin.txt", b"caf\xe9")) == "caf\ufffd"


def test_extract_text_strips_nul_characters(tmp_path):
    assert extract_text(write(tmp_path, "nul.txt", b"he\x00llo world")) == "hello world"


@pytest.mark.parametrize("content", [b"", b"  \n\t "], ids=["empty", "whitespace-only"])
def test_extract_text_rejects_files_without_text(tmp_path, content):
    with pytest.raises(ExtractionError, match="No extractable text found"):
        extract_text(write(tmp_path, "blank.txt", content))


def test_extract_text_reports_a_missing_file(tmp_path):
    with pytest.raises(ExtractionError, match="Stored file not found"):
        extract_text(tmp_path / "ghost.txt")


def test_extract_text_rejects_an_unsupported_extension(tmp_path):
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_text(write(tmp_path, "data.xyz", b"content"))


# ---------- extract_text: PDF ----------

def test_extract_text_reads_a_real_pdf(sample_pdf):
    text = extract_text(sample_pdf)

    assert "Sample PDF" in text
    assert "Lorem ipsum dolor sit amet" in text


def test_extract_text_joins_all_pages_of_a_pdf(tmp_path, sample_pdf):
    path = write(tmp_path, "two.pdf", two_page_pdf_bytes(sample_pdf))

    assert extract_text(path).count("Sample PDF") == 2


def test_extract_text_rejects_a_pdf_without_text(tmp_path):
    with pytest.raises(ExtractionError, match="No extractable text found"):
        extract_text(write(tmp_path, "blank.pdf", blank_pdf_bytes()))


@pytest.mark.parametrize("content", [b"this is not a pdf", b""], ids=["garbage", "zero-bytes"])
def test_extract_text_lets_unexpected_pdf_errors_propagate(tmp_path, content):
    with pytest.raises(Exception) as excinfo:
        extract_text(write(tmp_path, "bad.pdf", content))

    assert not isinstance(excinfo.value, ExtractionError)


# ---------- extract_text: DOCX ----------

def test_extract_text_reads_docx_paragraphs_in_order(tmp_path):
    path = write(tmp_path, "a.docx", docx_bytes(["First paragraph", "Second paragraph"]))

    assert extract_text(path).splitlines() == ["First paragraph", "Second paragraph"]


def test_extract_text_appends_docx_tables_after_the_body_and_skips_empty_cells(tmp_path):
    content = docx_bytes(
        ["Body one", "Body two"],
        table=[["A1", "B1"], ["A2", ""]],
    )

    text = extract_text(write(tmp_path, "t.docx", content))

    assert text.splitlines() == ["Body one", "Body two", "A1 | B1", "A2"]


def test_extract_text_rejects_an_empty_docx(tmp_path):
    with pytest.raises(ExtractionError, match="No extractable text found"):
        extract_text(write(tmp_path, "empty.docx", docx_bytes()))


def test_extract_text_lets_unexpected_docx_errors_propagate(tmp_path):
    with pytest.raises(Exception) as excinfo:
        extract_text(write(tmp_path, "bad.docx", b"not a zip"))

    assert not isinstance(excinfo.value, ExtractionError)


# ---------- process_document: the task itself ----------

def test_processing_a_text_file_completes_with_the_extracted_text(db, owner_id, tmp_path):
    doc_id = make_document(db, owner_id, stored_name=put_file(tmp_path, "doc.txt", b"hello world"))

    processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status == "completed"
    assert doc.processing_error is None
    assert doc.extracted_text == "hello world"


def test_status_is_processing_while_extraction_is_running(db, owner_id, tmp_path, monkeypatch):
    doc_id = make_document(db, owner_id, stored_name=put_file(tmp_path, "doc.txt", b"x"))
    observed = []

    def spy(path):
        with SessionLocal() as other_session:
            observed.append(other_session.get(Document, doc_id).processing_status)
        return "extracted"

    monkeypatch.setattr(processing, "extract_text", spy)

    processing.process_document(doc_id)

    assert observed == ["processing"]
    assert fresh(db, doc_id).processing_status == "completed"


@pytest.mark.parametrize(
    "name, content, expected",
    [
        ("empty.txt", b"", "No extractable text found"),
        ("blank.md", b"  \n\t ", "No extractable text found"),
        ("data.xyz", b"content", "Unsupported file type"),
    ],
    ids=["empty-file", "whitespace-only", "unsupported-extension"],
)
def test_known_problems_mark_the_document_failed_with_a_clear_message(
    db, owner_id, tmp_path, name, content, expected
):
    doc_id = make_document(db, owner_id, stored_name=put_file(tmp_path, name, content))

    processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status == "failed"
    assert expected in doc.processing_error
    assert doc.extracted_text is None


def test_a_missing_file_marks_the_document_failed(db, owner_id):
    doc_id = make_document(db, owner_id, stored_name="ghost.txt")

    processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status == "failed"
    assert doc.processing_error == "Stored file not found"


@pytest.mark.parametrize("name", ["bad.pdf", "bad.docx"])
def test_unexpected_extraction_errors_are_logged_but_hidden_from_the_user(
    db, owner_id, tmp_path, caplog, name
):
    doc_id = make_document(
        db, owner_id, stored_name=put_file(tmp_path, name, b"this is not a real document")
    )

    with caplog.at_level(logging.ERROR, logger="documind"):
        processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status == "failed"
    # python-docx's own error message contains the file path: it must never reach the user.
    assert doc.processing_error == GENERIC_FAILURE
    failures = [r for r in caplog.records if "Document processing failed" in r.getMessage()]
    assert len(failures) == 1
    assert failures[0].exc_info is not None


def test_processing_a_nonexistent_document_does_nothing():
    assert processing.process_document(999999) is None


def test_a_metadata_only_document_is_left_untouched(db, owner_id):
    doc_id = make_document(db, owner_id, stored_name=None, status=None)

    processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status is None
    assert doc.extracted_text is None


def test_text_with_nul_characters_completes_instead_of_failing(db, owner_id, tmp_path):
    stored = put_file(tmp_path, "nul.txt", b"he\x00llo world")
    doc_id = make_document(db, owner_id, stored_name=stored)

    processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status == "completed"
    assert doc.extracted_text == "hello world"


def test_reprocessing_a_failed_document_clears_the_old_error(db, owner_id, tmp_path):
    stored = put_file(tmp_path, "doc.txt", b"all good now")
    doc_id = make_document(db, owner_id, stored_name=stored, status="failed", error="old error")

    processing.process_document(doc_id)

    doc = fresh(db, doc_id)
    assert doc.processing_status == "completed"
    assert doc.processing_error is None


def test_an_unexpected_crash_is_logged_and_recorded_as_failed(
    db, owner_id, tmp_path, caplog, monkeypatch
):
    doc_id = make_document(db, owner_id, stored_name=put_file(tmp_path, "doc.txt", b"x"))
    monkeypatch.setattr(processing, "_run", explode)

    with caplog.at_level(logging.ERROR, logger="documind"):
        processing.process_document(doc_id)  # must not raise

    doc = fresh(db, doc_id)
    assert doc.processing_status == "failed"
    assert doc.processing_error == "Processing failed unexpectedly"
    assert any("Document processing crashed" in r.getMessage() for r in caplog.records)


def test_a_failure_while_recording_a_failure_is_logged_and_never_raised(
    db, owner_id, caplog, monkeypatch
):
    doc_id = make_document(db, owner_id, stored_name="whatever.txt")
    monkeypatch.setattr(processing, "_run", explode)
    monkeypatch.setattr(processing, "_mark_failed", explode)

    with caplog.at_level(logging.ERROR, logger="documind"):
        processing.process_document(doc_id)  # must not raise

    assert any("Could not record failure" in r.getMessage() for r in caplog.records)
    assert fresh(db, doc_id).processing_status == "pending"


def test_the_task_always_closes_its_session(monkeypatch):
    session = MagicMock()
    monkeypatch.setattr(processing, "SessionLocal", lambda: session)
    monkeypatch.setattr(processing, "_run", explode)
    monkeypatch.setattr(processing, "_mark_failed", explode)

    processing.process_document(1)

    session.close.assert_called_once()


def test_a_document_deleted_during_processing_is_handled_quietly(
    db, owner_id, tmp_path, monkeypatch
):
    doc_id = make_document(db, owner_id, stored_name=put_file(tmp_path, "doc.txt", b"x"))

    def delete_the_row_then_return_text(path):
        with SessionLocal() as other_session:
            other_session.query(Document).filter(Document.id == doc_id).delete()
            other_session.commit()
        return "extracted text"

    monkeypatch.setattr(processing, "extract_text", delete_the_row_then_return_text)

    processing.process_document(doc_id)  # must not raise

    assert db.query(Document).count() == 0


# ---------- end to end through the API ----------

def test_uploaded_sample_pdf_is_extracted(client, auth_headers, db, sample_pdf):
    uploaded = upload(client, auth_headers, "sample.pdf", sample_pdf.read_bytes()).json()
    assert uploaded["processing_status"] == "pending"

    fetched = client.get(f"{DOCUMENTS}/{uploaded['id']}", headers=auth_headers).json()

    assert fetched["processing_status"] == "completed"
    text = fresh(db, uploaded["id"]).extracted_text
    assert "Sample PDF" in text
    assert "Lorem ipsum dolor sit amet" in text


def test_uploaded_docx_includes_table_text(client, auth_headers, db):
    content = docx_bytes(["Quarterly report"], table=[["Region", "Revenue"], ["North", "100"]])
    uploaded = upload(client, auth_headers, "report.docx", content).json()

    fetched = client.get(f"{DOCUMENTS}/{uploaded['id']}", headers=auth_headers).json()

    assert fetched["processing_status"] == "completed"
    text = fresh(db, uploaded["id"]).extracted_text
    assert "Quarterly report" in text
    assert "Region | Revenue" in text
    assert "North | 100" in text


def test_uploaded_corrupt_pdf_fails_with_a_generic_message(client, auth_headers):
    response = upload(client, auth_headers, "bad.pdf", b"this is not a pdf")
    assert response.status_code == 201

    fetched = client.get(f"{DOCUMENTS}/{response.json()['id']}", headers=auth_headers).json()

    assert fetched["processing_status"] == "failed"
    assert fetched["processing_error"] == GENERIC_FAILURE


def test_uploaded_pdf_without_text_fails_with_a_clear_message(client, auth_headers):
    uploaded = upload(client, auth_headers, "scan.pdf", blank_pdf_bytes()).json()

    fetched = client.get(f"{DOCUMENTS}/{uploaded['id']}", headers=auth_headers).json()

    assert fetched["processing_status"] == "failed"
    assert "No extractable text" in fetched["processing_error"]


def test_a_failed_document_keeps_its_file_and_stays_downloadable(client, auth_headers):
    content = blank_pdf_bytes()
    uploaded = upload(client, auth_headers, "scan.pdf", content).json()

    download = client.get(f"{DOCUMENTS}/{uploaded['id']}/download", headers=auth_headers)

    assert download.status_code == 200
    assert download.content == content
    assert len(client.get(DOCUMENTS, headers=auth_headers).json()) == 1


def test_list_shows_the_processing_status_of_each_document(client, auth_headers):
    upload(client, auth_headers, "good.txt", b"fine")
    upload(client, auth_headers, "empty.txt", b"")

    listed = client.get(DOCUMENTS, headers=auth_headers).json()

    assert [d["processing_status"] for d in listed] == ["completed", "failed"]